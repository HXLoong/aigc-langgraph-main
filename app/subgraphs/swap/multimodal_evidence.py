"""Attachment locations and final authoritative bindings for swap orders."""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from typing import Any

import openpyxl
from pydantic import BaseModel, ConfigDict, Field

from app.extraction.fields import FieldCandidate, FieldRecord
from app.subgraphs.swap.models import SwapPlaceOrderParams
from app.subgraphs.swap.normalize import (
    currency,
    normalize_candidates,
    normalize_field,
    quantity_unit,
)


class ImageTranscription(BaseModel):
    """A transcription is model evidence, not proof against the original pixels."""

    model_config = ConfigDict(extra="forbid")
    text: str = Field(description="本张图片的逐字转写；保留原始行列、表头、单位和符号；不补全或推断")


@dataclass(frozen=True)
class ExcelEvidenceRow:
    reference: str
    values: dict[str, Any]
    sources: dict[str, str]


def excel_evidence_rows(content: bytes, file_index: int) -> list[ExcelEvidenceRow]:
    """Keep sheet/row/column locations; formulas without cached values fail explicitly."""
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    cached = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    result: list[ExcelEvidenceRow] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            rows = iter(sheet.iter_rows())
            headers_row = next(rows, ())
            headers = [str(cell.value).strip() if cell.value is not None else ""
                       for cell in headers_row]
            headers = ["交易对手" if header == "产品" else header for header in headers]
            nonempty = [header for header in headers if header]
            if len(set(nonempty)) != len(nonempty):
                raise ValueError("Excel 表头重复，无法确定字段归属")
            for row in rows:
                if all(cell.value is None for cell in row):
                    continue
                row_number = next(cell.row for cell in row if cell.value is not None)
                reference = f"file:{file_index}:sheet:{sheet_index}:row:{row_number}"
                values, sources = {}, {}
                for header, cell in zip(headers, row, strict=False):
                    value = cell.value
                    if cell.data_type == "f":
                        value = cached.worksheets[sheet_index].cell(cell.row, cell.column).value
                        if value is None:
                            raise ValueError("Excel 公式缺少缓存结果，请先计算并保存文件")
                    if value is None:
                        continue
                    if not header:
                        raise ValueError("Excel 数据列缺少表头，无法确定字段归属")
                    values[header] = value
                    sources[f"{reference}:column:{cell.column_letter}"] = f"{header}: {value}"
                sources[reference] = json.dumps(values, ensure_ascii=False, default=str)
                result.append(ExcelEvidenceRow(reference, values, sources))
    finally:
        workbook.close()
        cached.close()
    return result


def lock_attachment_bindings(
    params: SwapPlaceOrderParams,
    records: dict[str, FieldRecord],
    counterparties: list[dict[str, Any]],
) -> tuple[SwapPlaceOrderParams, dict[str, FieldRecord]]:
    """Keep attachment provenance; Java resolves the extracted security expression."""
    result = dict(records)
    rows = []
    for index, item in enumerate(params.order_list):
        row = item.model_dump()
        prefix = f"swap/place_order.orderList.{index}."
        if item.place_order_shortname:
            matches = {
                counterparty["shortName"]
                for counterparty in counterparties
                if counterparty.get("shortName") and item.place_order_shortname in {
                    counterparty.get("shortName"), counterparty.get("longName"),
                }
            }
            if len(matches) != 1:
                raise ValueError("附件交易对手未匹配唯一授权账户，请使用完整账户名称")
            key = prefix + "placeOrderShortname"
            result[key + ".candidate"] = result[key]
            shortname = next(iter(matches))
            row["placeOrderShortname"] = shortname
            result[key] = FieldRecord(
                value=shortname, source="goats", evidence=shortname,
                origin="authorized-counterparties", locked=True,
            )
        rows.append(row)
    return SwapPlaceOrderParams.model_validate({"orderList": rows}), result


def normalize_attachment_candidates(
    candidates: BaseModel, sources: dict[str, str],
) -> tuple[SwapPlaceOrderParams, dict[str, FieldRecord]]:
    """Apply units from an Excel column header in Code without inventing model evidence."""
    data = candidates.model_dump(by_alias=True)
    header_fields: list[tuple[int, str, FieldCandidate, str, str]] = []
    for index, row in enumerate(data.get("orderList") or []):
        for key in ("placeOrderQuantity", "placeOrderNotional"):
            cell = row.get(key)
            if not cell or cell.get("origin") != "attachment":
                continue
            value = cell.get("value") or ""
            if not re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", value):
                continue
            reference = cell.get("reference") or ""
            if ":column:" not in reference:
                # A row-level candidate still has to inherit its column's declared unit.
                prefixes = ("数量", "委托量") if key == "placeOrderQuantity" else ("金额", "名义本金")
                matches = [name.removeprefix("attachment:") for name, text in sources.items()
                           if name.startswith("attachment:" + reference + ":column:")
                           and text.partition(": ")[2] == value
                           and text.startswith(prefixes)]
                if len(matches) > 1:
                    raise ValueError("Excel 数值对应多个列，请明确单元格来源")
                if not matches:
                    continue
                reference = matches[0]
            source = sources.get("attachment:" + reference, "")
            header = source.split(": ", 1)[0]
            unit_match = re.search(r"[（(]([千万亿]?(?:股|手|元|美元|港元|人民币))[）)]$", header)
            if not unit_match:
                continue
            candidate = FieldCandidate.model_validate(cell).model_copy(update={"reference": reference})
            header_fields.append((index, key, candidate, unit_match[1], source))
            row[key] = None
    params, records = normalize_candidates(type(candidates).model_validate(data), sources)
    orders = [item.model_dump() for item in params.order_list]
    for index, key, candidate, suffix, source in header_fields:
        token = (candidate.value or "") + suffix
        unit = quantity_unit(token)
        target = "placeOrderNotional" if unit == "AMOUNT" else key
        value = normalize_field(target, token, source)
        orders[index][target] = value
        prefix = f"swap/place_order.orderList.{index}."
        record = FieldRecord(
            value=value, source="inferred", evidence=source,
            origin="attachment:" + (candidate.reference or ""),
            confidence=candidate.confidence, locked=True,
        )
        records[prefix + target] = record
        for derived_key, derived_value in (
            ("placeOrderQuantityUnit", unit),
            ("placeOrderNotionalCurrency", currency(token) if unit == "AMOUNT" else None),
        ):
            if derived_value is None:
                continue
            existing = orders[index].get(derived_key)
            if existing is not None and existing != derived_value:
                raise ValueError("Excel 单位与列标题冲突")
            orders[index][derived_key] = derived_value
            records[prefix + derived_key] = record.model_copy(update={"value": derived_value})
    return SwapPlaceOrderParams.model_validate({"orderList": orders}), records
