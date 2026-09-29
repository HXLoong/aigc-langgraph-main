"""Excel 确定性读取：保留单元格值、表头与位置，不解释业务语义。"""
from __future__ import annotations

import io
import json
from dataclasses import dataclass

import httpx
import openpyxl

from app.config import get_settings


@dataclass(frozen=True)
class ExcelInput:
    """所有数据行及单元格的原始证据集合。"""

    rows: list[str]
    sources: dict[str, str]


async def fetch_file(url: str) -> bytes:
    """下载附件只使用文件 URL，不携带 Java 或模型服务的鉴权头。"""
    timeout = get_settings().multimodal_fetch_timeout_seconds
    async with httpx.AsyncClient(timeout=timeout) as client:
        response = await client.get(url)
        response.raise_for_status()
        return response.content


def read_excel(content: bytes, file_index: int) -> ExcelInput:
    """读全部工作表，公式必须存在缓存值，表头和单位原样保留。"""
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    cached = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sources: dict[str, str] = {}
    rows: list[str] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            iterator = iter(sheet.iter_rows())
            headers = [str(cell.value) if cell.value is not None else "" for cell in next(iterator, ())]
            nonempty = [header for header in headers if header]
            if len(nonempty) != len(set(nonempty)):
                raise ValueError("Excel 表头重复，请明确字段名称")
            for cells in iterator:
                if all(cell.value is None for cell in cells):
                    continue
                row_number = next(cell.row for cell in cells if cell.value is not None)
                reference = f"file:{file_index}:sheet:{sheet_index}:row:{row_number}"
                values = {}
                for header, cell in zip(headers, cells, strict=True):
                    value = cell.value
                    if cell.data_type == "f":
                        value = cached.worksheets[sheet_index].cell(cell.row, cell.column).value
                        if value is None:
                            raise ValueError("Excel 公式缺少缓存结果，请计算并保存文件")
                    if value is None:
                        continue
                    if not header:
                        raise ValueError("Excel 数据列缺少表头")
                    values[header] = value
                    sources[f"{reference}:column:{cell.column_letter}"] = f"{header}: {value}"
                sources[reference] = json.dumps(values, ensure_ascii=False, default=str)
                rows.append(reference)
    finally:
        # 只读工作簿也持有压缩文件句柄，所有退出路径都显式释放。
        workbook.close()
        cached.close()
    return ExcelInput(rows, sources)
