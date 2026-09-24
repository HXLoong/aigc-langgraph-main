"""默认 Excel 导出必须完整保留现役业务集，避免读取落后的镜像。"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def test_default_export_contains_every_current_case(tmp_path: Path) -> None:
    target = tmp_path / "golden.xlsx"
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts/convert_jsonl_to_excel.py"), "--output", str(target)],
        cwd=tmp_path, capture_output=True, text=True, check=False,
    )
    assert result.returncode == 0, result.stderr
    expected = {
        row["caseNo"]: row
        for path in (ROOT / "tests/fixtures/biz").glob("*.jsonl")
        for line in path.read_text().splitlines() if line.strip()
        for row in [json.loads(line)]
    }
    workbook = load_workbook(target, read_only=True)
    try:
        actual = {}
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows)
            for cells in rows:
                row = dict(zip(headers, cells, strict=True))
                case_id = row["caseNo"]
                assert case_id not in actual
                actual[case_id] = row
        assert actual.keys() == expected.keys()
        for case_id, case in expected.items():
            assert actual[case_id]["send_text"] == case["send_text"]
            if case.get("sub_scenes"):
                assert json.loads(actual[case_id]["sub_scenes"]) == case["sub_scenes"]
    finally:
        workbook.close()
