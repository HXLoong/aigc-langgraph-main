"""CSV → Excel 的数据保真、CLI、格式及真实测试集链路。"""

from __future__ import annotations

import csv
import importlib
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "scripts" / "convert_csv_to_excel.py"


def run_cli(*args: str | Path, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *map(str, args)],
        cwd=cwd or ROOT,
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )


def write_csv(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        csv.writer(stream).writerows(rows)


def converter():
    assert SCRIPT.exists(), "CSV → Excel 脚本尚未实现"
    return importlib.import_module("scripts.convert_csv_to_excel")


def test_single_file_preserves_strings_and_applies_styles(tmp_path: Path) -> None:
    path = tmp_path / "中文.csv"
    rows = [
        ["用例ID", "操作步骤序号", "测试数据", "步骤级预期", "引用说明", "步骤场景"],
        ["001", "01", '  中文,"引用"\r\n下一行  ', '{"intent": "new_inquiry"}', "true", ""],
        ["12345678901234567890", "2", "=1+1", "#N/A", "false", "null"],
        ['=HYPERLINK("https://example.com")', "3", "+123", "-123", "@机器人", "正常\r末尾"],
        ["long", "4", "长文本" * 400, "", "", ""],
        ["", "", "", "", "", ""],
    ]
    write_csv(path, rows)
    result = run_cli("--input", path, cwd=tmp_path)
    assert result.returncode == 0, result.stderr
    workbook = openpyxl.load_workbook(path.with_suffix(".xlsx"))
    try:
        assert workbook.sheetnames == ["中文"]
        sheet = workbook.active
        assert [[c.value or "" for c in row] for row in sheet] == rows
        for row in sheet:
            for cell in row:
                if cell.value:
                    assert cell.data_type == "s"
                assert cell.font.name == "等线"
                assert cell.font.sz == 11
                assert cell.alignment.wrap_text is True
                assert cell.alignment.vertical == "top"
        assert sheet["A1"].fill.fgColor.rgb == "001F4E78"
        assert sheet["A1"].font.color.rgb == "00FFFFFF"
        assert sheet["A1"].font.bold
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:F6"
        assert len(sheet.merged_cells.ranges) == 0
        assert [sheet.column_dimensions[c].width for c in "ABCDEF"] == [28, 16, 60, 60, 40, 24]
        assert sheet.row_dimensions[1].height == 32
        assert all(30 <= sheet.row_dimensions[n].height <= 409 for n in range(2, 7))
        assert sheet.row_dimensions[5].height > sheet.row_dimensions[3].height
    finally:
        workbook.close()


def test_batch_names_order_dry_run_and_overwrite(tmp_path: Path) -> None:
    folder = tmp_path / "csv"
    folder.mkdir()
    (folder / "nested").mkdir()
    names = ["a[场景]", "a_场景_", "b" * 32 + "1", "b" * 32 + "2", "'quoted'"]
    for name in names:
        write_csv(folder / f"{name}.csv", [["数据"], [name]])
    write_csv(folder / "nested" / "ignored.csv", [["数据"], ["忽略"]])
    result = run_cli("--input", folder, "--dry-run")
    assert result.returncode == 0, result.stderr
    assert not (folder / "黄金数据集.xlsx").exists()
    assert "sheet=" in result.stderr and "rows=1" in result.stderr
    result = run_cli("--input", folder)
    assert result.returncode == 0, result.stderr
    target = folder / "黄金数据集.xlsx"
    workbook = openpyxl.load_workbook(target)
    try:
        assert len(workbook.sheetnames) == 5
        assert len({n.casefold() for n in workbook.sheetnames}) == 5
        assert all(
            len(n) <= 31 and not any(c in n for c in "[]:*?/\\") for n in workbook.sheetnames
        )
        assert all(not n.startswith("'") and not n.endswith("'") for n in workbook.sheetnames)
        assert [s["A2"].value for s in workbook] == sorted(names)
        assert "a_场景__2" in workbook.sheetnames
        assert "b" * 29 + "_2" in workbook.sheetnames
    finally:
        workbook.close()
    write_csv(folder / f"{names[0]}.csv", [["数据"], ["已更新"]])
    result = run_cli("--input", folder)
    assert result.returncode == 0, result.stderr
    workbook = openpyxl.load_workbook(target)
    try:
        assert workbook["a_场景_"]["A2"].value == "已更新"
    finally:
        workbook.close()
    assert not list(folder.glob("*.tmp"))


def test_sheet_names_handle_case_insensitive_collisions() -> None:
    module = converter()
    used: set[str] = set()
    assert [module.sheet_name(name, used) for name in ["Case", "case", "CASE", "[]", "''"]] == [
        "Case",
        "case_2",
        "CASE_3",
        "__",
        "Sheet",
    ]


def test_blank_lines_header_only_and_explicit_output(tmp_path: Path) -> None:
    path = tmp_path / "empty-data.csv"
    path.write_text("\n\n列1,列2\n\n", encoding="utf-8")
    target = tmp_path / "out" / "chosen.xlsx"
    result = run_cli("--input", path, "--output", target)
    assert result.returncode == 0, result.stderr
    workbook = openpyxl.load_workbook(target)
    try:
        assert workbook.active.max_row == 1
        assert workbook.active.max_column == 2
        assert workbook.active.auto_filter.ref == "A1:B1"
    finally:
        workbook.close()


def test_trailing_fully_empty_csv_columns_are_not_exported(tmp_path: Path) -> None:
    path = tmp_path / "trailing-empty.csv"
    write_csv(path, [["A", "B", "", ""], ["x", "y", "", ""]])

    result = run_cli("--input", path)

    assert result.returncode == 0, result.stderr
    workbook = openpyxl.load_workbook(path.with_suffix(".xlsx"))
    try:
        assert workbook.active.max_column == 2
        assert [cell.value for cell in workbook.active[1]] == ["A", "B"]
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("content", "message"),
    [
        ("\n\n", "表头"),
        (",B\nx,y\n", "表头"),
        ("A,A\nx,y\n", "重复"),
        ("A,B\nx\n", "字段数"),
        ("A\nx,y\n", "字段数"),
        ('A\n"unclosed', "CSV"),
        ("A\n非法\x00字符\n", "非法字符"),
        ("A\n非法\ufffe字符\n", "非法字符"),
        ("A\n" + "x" * 32768 + "\n", "32767"),
        ("A\n" + "😀" * 16384 + "\n", "32767"),
    ],
    ids=[
        "empty",
        "empty-header",
        "duplicate-header",
        "short-row",
        "long-row",
        "bad-quote",
        "nul",
        "noncharacter",
        "long-cell",
        "long-utf16-cell",
    ],
)
def test_invalid_batch_preserves_existing_output(
    tmp_path: Path, content: str, message: str
) -> None:
    write_csv(tmp_path / "a.csv", [["A"], ["有效"]])
    invalid = tmp_path / "z.csv"
    invalid.write_text(content, encoding="utf-8")
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"original workbook")
    result = run_cli("--input", tmp_path, "--output", target)
    assert result.returncode != 0
    assert "z.csv" in result.stderr and message in result.stderr
    assert target.read_bytes() == b"original workbook"
    assert not list(tmp_path.glob("*.tmp"))


def test_bad_cell_error_reports_record_and_column(tmp_path: Path) -> None:
    path = tmp_path / "bad.csv"
    path.write_text('A,B\n"跨\n行",bad\x01text\n', encoding="utf-8")
    result = run_cli("--input", path)
    assert result.returncode != 0
    assert "记录 2" in result.stderr and "列 2" in result.stderr
    assert not path.with_suffix(".xlsx").exists()


def test_missing_inputs_and_wrong_output_extension(tmp_path: Path) -> None:
    for source in [tmp_path / "missing.csv", tmp_path]:
        result = run_cli("--input", source)
        assert result.returncode != 0
        assert "convert_jsonl_to_csv.py" in result.stderr
    path = tmp_path / "a.csv"
    write_csv(path, [["A"], ["x"]])
    original = path.read_bytes()
    result = run_cli("--input", path, "--output", path)
    assert result.returncode != 0 and ".xlsx" in result.stderr
    assert path.read_bytes() == original


def test_excel_dimension_limits_are_validated(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    module = converter()
    path = tmp_path / "limits.csv"
    write_csv(path, [["A", "B"], ["x", "y"]])
    monkeypatch.setattr(module, "MAX_COLUMNS", 1)
    assert module.main(["--input", str(path), "--dry-run"]) == 1
    monkeypatch.setattr(module, "MAX_COLUMNS", 16384)
    monkeypatch.setattr(module, "MAX_ROWS", 1)
    assert module.main(["--input", str(path), "--dry-run"]) == 1
    assert not path.with_suffix(".xlsx").exists()


@pytest.mark.parametrize("failure", ["save", "replace"])
def test_save_failure_preserves_destination_and_cleans_temp(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    failure: str,
) -> None:
    module = converter()
    path = tmp_path / "data.csv"
    write_csv(path, [["A"], ["x"]])
    target = path.with_suffix(".xlsx")
    target.write_bytes(b"original")

    def fail(*args, **kwargs):
        raise OSError("模拟文件被占用")

    if failure == "save":
        monkeypatch.setattr(module.Workbook, "save", fail)
    else:
        monkeypatch.setattr(module.os, "replace", fail)
    assert module.main(["--input", str(path)]) == 1
    assert target.read_bytes() == b"original"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["data.csv", "data.xlsx"]


def test_default_input_is_relative_to_repository(tmp_path: Path) -> None:
    result = run_cli("--dry-run", cwd=tmp_path)
    assert result.returncode == 0, result.stderr
    assert str(ROOT / "tests" / "fixtures" / "categories" / "csv") in result.stderr
    assert "files=7" in result.stderr
    assert not list(tmp_path.iterdir())


def test_real_jsonl_csv_excel_roundtrip(tmp_path: Path) -> None:
    csv_dir = tmp_path / "csv"
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "convert_jsonl_to_csv.py"),
            "--output-dir",
            str(csv_dir),
        ],
        capture_output=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    target = tmp_path / "result.xlsx"
    result = run_cli("--input", csv_dir, "--output", target)
    assert result.returncode == 0, result.stderr
    workbook = openpyxl.load_workbook(target)
    try:
        csv_paths = sorted(csv_dir.glob("*.csv"))
        assert csv_paths
        assert len(workbook.worksheets) == len(csv_paths)
        for sheet, csv_path in zip(workbook, csv_paths, strict=True):
            assert csv_path.stem.startswith(sheet.title)
            with csv_path.open(encoding="utf-8-sig", newline="") as stream:
                expected = list(csv.reader(stream))
            assert sheet.max_row == len(expected)
            assert sheet.max_column == 22
            assert [[c.value or "" for c in row] for row in sheet] == expected
            assert sheet.freeze_panes == "A2"
            assert sheet.auto_filter.ref == f"A1:V{sheet.max_row}"
    finally:
        workbook.close()
