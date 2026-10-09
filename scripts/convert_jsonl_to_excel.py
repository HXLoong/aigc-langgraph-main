"""将 JSONL 测试集导出为一个 Excel 工作簿（先经 CSV 中转）。

用法：
    python scripts/convert_jsonl_to_excel.py
    python scripts/convert_jsonl_to_excel.py --input tests/datasets/biz
    python scripts/convert_jsonl_to_excel.py --input example.jsonl --output example.xlsx

默认读取仓库 tests/datasets/biz，输出 tests/datasets/biz/黄金数据集.xlsx。
流程：先调用 convert_jsonl_to_csv.py 把 JSONL 转成逐步中文 CSV，再把每个 CSV
写成一张 sheet。输出 Excel 已存在时**追加新 sheet**（同名自动加 _2 等后缀），
保留已有 sheet 不覆盖；显式 --output 优先。

CSV → Excel 的样式、校验与原子替换能力原在 convert_csv_to_excel.py，现并入本脚本。
表头及所有非空内容按文本保留，不解析 JSON、不转换数字或日期、不执行公式。
每张表冻结首行、开启筛选与自动换行；依赖项目已有的 openpyxl。
"""

from __future__ import annotations

import argparse
import csv
import logging
import math
import os
import re
import tempfile
import unicodedata
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import convert_jsonl_to_csv as jsonl_converter
except ImportError:  # 以 scripts.convert_jsonl_to_excel 导入时
    from scripts import convert_jsonl_to_csv as jsonl_converter

logger = logging.getLogger(__name__)
REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = REPO_ROOT / "tests" / "datasets" / "biz"
MAX_ROWS = 1_048_576
MAX_COLUMNS = 16_384
MAX_CELL_LENGTH = 32_767
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
INVALID_SHEET = re.compile(r"[\\/*?:\[\]\x00-\x1f\ufffe\uffff]")
TEXT_COLUMNS = {
    "测试数据",
    "用例级预期",
    "步骤级预期",
    "响应必须包含",
    "响应包含任一",
    "响应不得包含",
}
DETAIL_COLUMNS = {"引用说明", "用例其他字段", "步骤其他字段"}
ID_COLUMNS = {"来源文件", "用例ID", "用例编号", "用例名称", "数据来源"}
SHORT_COLUMNS = {"源行号", "操作步骤序号", "是否@机器人", "是否引用上一条"}


@dataclass
class CsvTable:
    path: Path
    title: str
    rows: list[list[str]]  # 包含表头，空字段仍占据对应列。


def sheet_name(stem: str, used: set[str]) -> str:
    """生成 Excel 合法标题；used 保存已使用标题的 casefold 值。"""
    base = INVALID_SHEET.sub("_", stem).strip("'") or "Sheet"
    title = base[:31].rstrip("'") or "Sheet"
    number = 2
    while title.casefold() in used:
        suffix = f"_{number}"
        title = base[: 31 - len(suffix)] + suffix
        number += 1
    used.add(title.casefold())
    return title


def validate_cell(value: str, location: str) -> None:
    invalid = INVALID_XML.search(value)
    if invalid:
        raise ValueError(f"{location}: Excel/XML 非法字符 U+{ord(invalid[0]):04X}")
    # Excel 使用 UTF-16，非 BMP 字符占两个编码单元；提前拒绝，避免写入时截断。
    if len(value.encode("utf-16-le")) // 2 > MAX_CELL_LENGTH:
        raise ValueError(f"{location}: 单元格超过 {MAX_CELL_LENGTH} 字符限制")


def read_csv(path: Path, title: str) -> CsvTable:
    rows: list[list[str]] = []
    previous_limit = csv.field_size_limit()
    csv.field_size_limit(max(previous_limit, MAX_CELL_LENGTH * 4))
    record_number = 0
    try:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream, strict=True)
            for row in reader:
                if not row:
                    continue
                record_number += 1
                location = f"{path}: 记录 {record_number} (截至文件行 {reader.line_num})"
                if record_number > MAX_ROWS:
                    raise ValueError(f"{location}: 超过 Excel 行数限制 {MAX_ROWS}（含表头）")
                if len(row) > MAX_COLUMNS:
                    raise ValueError(f"{location}: 超过 Excel 列数限制 {MAX_COLUMNS}")
                if rows and len(row) != len(rows[0]):
                    raise ValueError(f"{location}: 字段数 {len(row)}，表头字段数 {len(rows[0])}")
                for column, value in enumerate(row, 1):
                    validate_cell(value, f"{location}, 列 {column}")
                rows.append(row)
    except (csv.Error, UnicodeError) as exc:
        raise ValueError(f"{path}: 记录 {record_number + 1}: CSV 读取失败: {exc}") from exc
    finally:
        csv.field_size_limit(previous_limit)
    if not rows:
        raise ValueError(f"{path}: CSV 为空，缺少表头")
    # 旧版导出曾在行尾追加若干完全空白列，统一裁掉；任一数据行有值的空表头仍拒绝。
    column_count = len(rows[0])
    while column_count > 0 and not rows[0][column_count - 1].strip() and all(
        not row[column_count - 1].strip() for row in rows[1:]
    ):
        column_count -= 1
    rows = [row[:column_count] for row in rows]
    if not rows[0]:
        raise ValueError(f"{path}: 记录 1, 列 1: 表头不能为空")
    seen: set[str] = set()
    for column, value in enumerate(rows[0], 1):
        if not value.strip():
            raise ValueError(f"{path}: 记录 1, 列 {column}: 表头不能为空")
        if value in seen:
            raise ValueError(f"{path}: 记录 1, 列 {column}: 重复表头 {value!r}")
        seen.add(value)
    return CsvTable(path, title, rows)


def column_width(header: str) -> int:
    if header in TEXT_COLUMNS:
        return 60
    if header in DETAIL_COLUMNS:
        return 40
    if header in ID_COLUMNS:
        return 28
    if header in SHORT_COLUMNS:
        return 16
    return 24


def wrapped_lines(value: str, width: int) -> int:
    """以中文全宽字符计两格，保守估算换行；仅用于显示，不改变原始值。"""
    lines = value.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    count = 0
    for line in lines:
        display_width = sum(
            4
            if char == "\t"
            else 0
            if unicodedata.combining(char)
            else 2
            if unicodedata.east_asian_width(char) in {"W", "F"}
            else 1
            for char in line
        )
        count += max(1, math.ceil(display_width / max(1, width - 2)))
    return count


def write_sheet(workbook: Workbook, table: CsvTable) -> None:
    sheet = workbook.create_sheet(table.title)
    widths = [column_width(header) for header in table.rows[0]]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    body_font = Font(name="等线", size=11)
    header_font = Font(name="等线", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    alignment = Alignment(vertical="top", wrap_text=True)
    for row_number, values in enumerate(table.rows, 1):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column)
            cell.value = value
            cell.data_type = "s"
            cell.number_format = "@"
            cell.font = header_font if row_number == 1 else body_font
            cell.alignment = alignment
            if row_number == 1:
                cell.fill = header_fill
        if row_number == 1:
            sheet.row_dimensions[row_number].height = 32
        else:
            line_count = max(
                wrapped_lines(value, width) for value, width in zip(values, widths, strict=True)
            )
            sheet.row_dimensions[row_number].height = min(409, max(30, line_count * 16 + 6))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{len(table.rows)}"


def save_workbook(workbook: Workbook, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    repaired: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output.parent, prefix=f".{output.name}.", suffix=".tmp", delete=False
        ) as stream:
            temporary = Path(stream.name)
        workbook.save(temporary)
        # 没有 lxml 的环境会把原始 CR 直接写入 XML，按原单元格恢复 CR 并转字符引用。
        corrections = {
            f"xl/worksheets/sheet{index}.xml": {
                cell.coordinate: cell.value
                for row in sheet
                for cell in row
                if isinstance(cell.value, str) and "\r" in cell.value
            }
            for index, sheet in enumerate(workbook, 1)
        }
        corrections = {name: values for name, values in corrections.items() if values}
        if corrections:
            with tempfile.NamedTemporaryFile(
                dir=output.parent, prefix=f".{output.name}.xml.", suffix=".tmp", delete=False
            ) as stream:
                repaired = Path(stream.name)
            namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            with ZipFile(temporary) as source, ZipFile(repaired, "w") as target:
                for entry in source.infolist():
                    data = source.read(entry.filename)
                    if entry.filename in corrections:
                        root = ET.fromstring(data)
                        values = corrections[entry.filename]
                        for element in root.iter(f"{namespace}c"):
                            if element.get("r") in values:
                                text = element.find(f"{namespace}is/{namespace}t")
                                if text is None:
                                    raise ValueError("无法恢复单元格换行：缺少文本节点")
                                text.text = values[element.get("r")]
                        data = ET.tostring(root, encoding="utf-8").replace(b"\r", b"&#13;")
                    target.writestr(entry, data)
            os.replace(repaired, temporary)
        os.replace(temporary, output)
    finally:
        if repaired is not None:
            repaired.unlink(missing_ok=True)
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _jsonl_to_csv(source: Path, csv_dir: Path) -> None:
    code = jsonl_converter.main(["--input", str(source), "--output-dir", str(csv_dir)])
    if code != 0:
        raise ValueError(f"JSONL → CSV 转换失败: {source}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="JSONL 文件或目录")
    parser.add_argument("--output", type=Path, help="输出 .xlsx 路径")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    workbook: Workbook | None = None
    try:
        source = args.input.resolve()
        output = (
            args.output.resolve()
            if args.output is not None
            else (source / "黄金数据集.xlsx" if source.is_dir() else source.with_suffix(".xlsx"))
        )
        if output.suffix.lower() != ".xlsx":
            raise ValueError(f"输出文件必须为 .xlsx: {output}")
        with tempfile.TemporaryDirectory(prefix="jsonl-to-excel-") as tmp:
            csv_dir = Path(tmp)
            _jsonl_to_csv(source, csv_dir)
            paths = sorted(p for p in csv_dir.glob("*.csv") if p.is_file())
            if not paths:
                raise ValueError(f"目录中没有 CSV 文件: {csv_dir}")
            # 输出已存在 → 加载并追加新 sheet；否则新建工作簿。
            if output.exists():
                workbook = load_workbook(output)
                used = {sheet.title.casefold() for sheet in workbook.worksheets}
            else:
                workbook = Workbook()
                workbook.remove(workbook.active)
                used = set()
            tables: list[CsvTable] = []
            for path in paths:
                title = sheet_name(path.stem, used)
                tables.append(read_csv(path, title))
            for table in tables:
                logger.info(
                    "input=%s sheet=%s rows=%s columns=%s",
                    table.path,
                    table.title,
                    len(table.rows) - 1,
                    len(table.rows[0]),
                )
                write_sheet(workbook, table)
            save_workbook(workbook, output)
            logger.info(
                "files=%s rows=%s output=%s",
                len(tables),
                sum(len(table.rows) - 1 for table in tables),
                output,
            )
    except (OSError, ValueError) as exc:
        logger.error("conversion_failed error=%s", exc)
        return 1
    finally:
        if workbook is not None:
            workbook.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
