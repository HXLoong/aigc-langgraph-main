"""将 Excel 工作簿还原为 JSONL（先抽取 sheet 为 CSV，再调用 convert_csv_to_jsonl.py）。

用法：
    python scripts/convert_excel_to_jsonl.py --input 黄金数据集.xlsx
    python scripts/convert_excel_to_jsonl.py --input 黄金数据集.xlsx --output-dir out

每个 sheet 是一张 convert_jsonl_to_csv.py 的 18 列中文 CSV 表。流程：先把每个
sheet 抽成 CSV（命名取自 sheet 首行"来源文件"列，缺省回退 sheet 名），再调用
convert_csv_to_jsonl.py 还原为 JSONL。空 sheet 跳过。

默认输出到输入文件旁的 <文件名>_jsonl/。整批校验通过后才写文件。中间 CSV
写入临时目录，进程结束即清理。
"""

from __future__ import annotations

import argparse
import csv
import logging
import re
import tempfile
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

try:
    import convert_csv_to_jsonl as csv_module
except ImportError:  # 以 scripts.convert_excel_to_jsonl 导入时
    from scripts import convert_csv_to_jsonl as csv_module

logger = logging.getLogger(__name__)
SOURCE_COLUMN = "来源文件"
_SAFE_JSONL = re.compile(r"^[A-Za-z0-9_.-]+\.jsonl$")


def _safe_stem(value: str) -> str:
    """sheet 标题 → 安全的 CSV 文件名 stem。"""
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .") or "Sheet"
    return stem


def _source_stem(header: list[str], data: list[list[Any]]) -> str | None:
    """从"来源文件"列取首个非空 .jsonl 文件名，返回其 stem；不可用返回 None。"""
    if SOURCE_COLUMN not in header:
        return None
    index = header.index(SOURCE_COLUMN)
    for row in data:
        if index >= len(row):
            continue
        value = row[index]
        if value is None:
            continue
        name = str(value).strip()
        if _SAFE_JSONL.fullmatch(name):
            return name[:-6]  # 去掉 .jsonl
        return None
    return None


def _sheet_to_csv(sheet, csv_dir: Path) -> Path | None:
    rows = [
        list(cells)
        for cells in sheet.iter_rows(values_only=True)
        if any(cell is not None and str(cell).strip() for cell in cells)
    ]
    if not rows:
        return None
    header = ["" if cell is None else str(cell) for cell in rows[0]]
    data = rows[1:]
    stem = _source_stem(header, data) or _safe_stem(sheet.title)
    path = csv_dir / f"{stem}.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        for row in data:
            writer.writerow(["" if cell is None else str(cell) for cell in row])
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--input", type=Path, required=True, help="输入 .xlsx 文件")
    parser.add_argument("--output-dir", type=Path, help="默认：输入文件旁的 <文件名>_jsonl 目录")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    try:
        source = args.input.resolve()
        if source.suffix.lower() != ".xlsx" or not source.is_file():
            raise ValueError(f"输入必须是存在的 .xlsx 文件: {source}")
        output_dir = (
            args.output_dir.resolve()
            if args.output_dir is not None
            else source.parent / f"{source.stem}_jsonl"
        )
        with tempfile.TemporaryDirectory(prefix="excel-to-jsonl-") as tmp:
            csv_dir = Path(tmp)
            workbook = load_workbook(source, read_only=True, data_only=False)
            try:
                csv_paths = [
                    path
                    for sheet in workbook
                    for path in [_sheet_to_csv(sheet, csv_dir)]
                    if path is not None
                ]
            finally:
                workbook.close()
            if not csv_paths:
                raise ValueError(f"{source}: 没有可转换的 sheet")
            code = csv_module.main(["--input", str(csv_dir), "--output-dir", str(output_dir)])
            if code != 0:
                raise ValueError("CSV → JSONL 转换失败")
    except (OSError, ValueError, BadZipFile, InvalidFileException) as exc:
        logger.error("conversion_failed error=%s", exc)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
